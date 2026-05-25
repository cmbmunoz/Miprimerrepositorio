"""Generate a final invoice Excel file from approved OC data."""

import os
import logging
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

log = logging.getLogger(__name__)

# ── Palette ───────────────────────────────────────────────────────────────────
DARK_GREEN  = "1D5C2E"
MED_GREEN   = "2E7D47"
LIGHT_GREEN = "D4EDDA"
WHITE       = "FFFFFF"
GRAY        = "808080"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def _font(bold=False, size=10, color="000000", italic=False) -> Font:
    return Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)

def _border() -> Border:
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center")

def _right() -> Alignment:
    return Alignment(horizontal="right", vertical="center")

def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center")


def generate_factura(oc_data: dict, output_dir: str, numero_factura: str) -> str:
    moneda = oc_data.get("moneda", "ARS")
    num_fmt = f'"{moneda} "#,##0.00'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Factura"

    for col, w in zip("ABCDE", (5, 38, 13, 18, 18)):
        ws.column_dimensions[col].width = w

    r = 1

    # ── Title ─────────────────────────────────────────────────────────────────
    ws.merge_cells(f"A{r}:E{r}")
    ws[f"A{r}"] = "FACTURA"
    ws[f"A{r}"].font = _font(bold=True, size=20, color=DARK_GREEN)
    ws[f"A{r}"].alignment = _center()
    ws.row_dimensions[r].height = 44
    r += 1

    # ── Metadata ──────────────────────────────────────────────────────────────
    meta = [
        ("N° Factura",  numero_factura,                   "Fecha emisión", datetime.now().strftime("%d/%m/%Y")),
        ("Ref. OC",     oc_data.get("numero_oc") or "",   "Fecha OC",      oc_data.get("fecha") or ""),
    ]
    for label_a, val_a, label_d, val_d in meta:
        ws[f"A{r}"] = label_a; ws[f"A{r}"].font = _font(bold=True)
        ws[f"B{r}"] = val_a;   ws[f"B{r}"].font = _font()
        ws[f"D{r}"] = label_d; ws[f"D{r}"].font = _font(bold=True)
        ws[f"E{r}"] = val_d;   ws[f"E{r}"].font = _font()
        ws.row_dimensions[r].height = 18
        r += 1

    r += 1

    # ── Proveedor ─────────────────────────────────────────────────────────────
    ws.merge_cells(f"A{r}:E{r}")
    ws[f"A{r}"] = "EMISOR (PROVEEDOR)"
    ws[f"A{r}"].font = _font(bold=True, color=WHITE)
    ws[f"A{r}"].fill = _fill(DARK_GREEN)
    ws[f"A{r}"].alignment = _left()
    ws.row_dimensions[r].height = 22
    r += 1

    proveedor = oc_data.get("proveedor", {})
    prov_rows = [
        ("Razón Social", proveedor.get("nombre", ""), "CUIT", proveedor.get("cuit", "")),
    ]
    for label_a, val_a, label_d, val_d in prov_rows:
        ws[f"A{r}"] = label_a; ws[f"A{r}"].font = _font(bold=True)
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = val_a;   ws[f"B{r}"].font = _font()
        ws[f"D{r}"] = label_d; ws[f"D{r}"].font = _font(bold=True)
        ws[f"E{r}"] = val_d;   ws[f"E{r}"].font = _font()
        ws.row_dimensions[r].height = 18
        r += 1

    r += 1

    # ── Cliente ───────────────────────────────────────────────────────────────
    ws.merge_cells(f"A{r}:E{r}")
    ws[f"A{r}"] = "RECEPTOR (CLIENTE)"
    ws[f"A{r}"].font = _font(bold=True, color=WHITE)
    ws[f"A{r}"].fill = _fill(MED_GREEN)
    ws[f"A{r}"].alignment = _left()
    ws.row_dimensions[r].height = 22
    r += 1

    cliente = oc_data.get("cliente", {})
    cli_rows = [
        ("Razón Social", cliente.get("nombre", ""),   "CUIT",      cliente.get("cuit", "")),
        ("Dirección",    cliente.get("direccion", ""), "",          ""),
    ]
    for label_a, val_a, label_d, val_d in cli_rows:
        ws[f"A{r}"] = label_a; ws[f"A{r}"].font = _font(bold=True)
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = val_a;   ws[f"B{r}"].font = _font()
        ws[f"D{r}"] = label_d; ws[f"D{r}"].font = _font(bold=True)
        ws[f"E{r}"] = val_d;   ws[f"E{r}"].font = _font()
        ws.row_dimensions[r].height = 18
        r += 1

    r += 1

    # ── Items ─────────────────────────────────────────────────────────────────
    headers = ("#", "Descripción", "Cantidad", "Precio Unit.", "Subtotal")
    aligns  = (_center(), _center(), _center(), _right(), _right())
    for col, (hdr, aln) in zip("ABCDE", zip(headers, aligns)):
        c = ws[f"{col}{r}"]
        c.value = hdr
        c.font = _font(bold=True, color=WHITE)
        c.fill = _fill(MED_GREEN)
        c.alignment = aln
        c.border = _border()
    ws.row_dimensions[r].height = 24
    r += 1

    items = oc_data.get("items", [])
    for i, item in enumerate(items, start=1):
        row_fill = _fill(LIGHT_GREEN) if i % 2 == 0 else None
        values = {
            "A": i,
            "B": item.get("descripcion", ""),
            "C": item.get("cantidad", 0),
            "D": item.get("precio_unitario", 0.0),
            "E": item.get("subtotal", 0.0),
        }
        aligns_row = {"A": _center(), "B": _left(), "C": _center(), "D": _right(), "E": _right()}
        for col, val in values.items():
            c = ws[f"{col}{r}"]
            c.value = val
            c.font = _font()
            c.border = _border()
            c.alignment = aligns_row[col]
            if row_fill:
                c.fill = row_fill
            if col in ("D", "E"):
                c.number_format = num_fmt
        ws.row_dimensions[r].height = 18
        r += 1

    r += 1

    # ── Totals ────────────────────────────────────────────────────────────────
    iva_pct = oc_data.get("iva_porcentaje", 21)
    totals = [
        (f"Subtotal ({moneda})",  oc_data.get("subtotal", 0.0),  False),
        (f"IVA {iva_pct}%",       oc_data.get("iva_monto", 0.0), False),
        (f"TOTAL {moneda}",       oc_data.get("total", 0.0),      True),
    ]
    for label, value, is_total in totals:
        ws.merge_cells(f"A{r}:C{r}")
        d = ws[f"D{r}"]
        e = ws[f"E{r}"]
        d.value = label
        e.value = value
        e.number_format = num_fmt
        d.border = _border()
        e.border = _border()
        d.alignment = _right()
        e.alignment = _right()
        if is_total:
            d.font = _font(bold=True, size=11, color=WHITE)
            d.fill = _fill(DARK_GREEN)
            e.font = _font(bold=True, size=11, color=WHITE)
            e.fill = _fill(DARK_GREEN)
        else:
            d.font = _font(bold=True)
            e.font = _font()
        ws.row_dimensions[r].height = 22
        r += 1

    # ── Notes ─────────────────────────────────────────────────────────────────
    notas = oc_data.get("notas")
    if notas:
        r += 1
        ws[f"A{r}"] = "Observaciones:"
        ws[f"A{r}"].font = _font(bold=True)
        r += 1
        ws.merge_cells(f"A{r}:E{r}")
        ws[f"A{r}"] = notas
        ws[f"A{r}"].font = _font(italic=True)
        ws[f"A{r}"].alignment = Alignment(wrap_text=True)
        ws.row_dimensions[r].height = 40

    # ── Footer ────────────────────────────────────────────────────────────────
    r += 2
    ws.merge_cells(f"A{r}:E{r}")
    ws[f"A{r}"] = f"Factura N° {numero_factura} — emitida el {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws[f"A{r}"].font = _font(italic=True, size=9, color=GRAY)
    ws[f"A{r}"].alignment = _center()

    # ── Save ──────────────────────────────────────────────────────────────────
    os.makedirs(output_dir, exist_ok=True)
    filename = f"factura_{numero_factura}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    log.info("Factura guardada: %s", filepath)
    return filepath
